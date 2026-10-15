import base64

import docx
from pypdf import PdfReader

from conftest import SAMPLE_PROFILE
from jobagent import onboarding, settings
from jobagent.profile import load_profile
from jobagent.tailor.render import render_html, render_pdf


def test_bullet_ids_are_positional():
    profile = load_profile(SAMPLE_PROFILE)
    ids = profile.bullets()
    assert ids["exp1.1"].startswith("Built a refund-reconciliation service")
    assert ids["exp2.3"].startswith("Wrote Celery pipelines")
    assert ids["proj1.1"].startswith("Parses bank SMS exports")
    assert len(ids) == 8
    assert "Kafka" in profile.all_skills()
    assert profile.is_placeholder()


def test_extract_profile_sends_pdf_and_parses(fake_api, tmp_path):
    fake_api.reply(SAMPLE_PROFILE.read_text())
    resume = tmp_path / "resume.pdf"
    resume.write_bytes(b"%PDF-1.4 fake")

    profile = onboarding.extract_profile(resume)

    assert profile.name == "Sample Candidate"
    content = fake_api.sent[0]["body"]["messages"][0]["content"]
    assert content[0]["type"] == "document"
    assert content[0]["source"]["media_type"] == "application/pdf"
    assert base64.b64decode(content[0]["source"]["data"]) == b"%PDF-1.4 fake"
    assert fake_api.sent[0]["body"]["output_config"]["effort"] == "medium"


def test_extract_profile_fills_link_targets_and_name_case(fake_api, tmp_path):
    # Claude reads text only: it sees "GitHub" but not the URL behind it, and copies "SAMPLE CANDIDATE" as printed.
    sample = load_profile(SAMPLE_PROFILE)
    resume = render_pdf(sample, tmp_path / "resume.pdf")  # real PDF whose links carry their targets
    as_claude_sees_it = sample.model_copy(deep=True)
    as_claude_sees_it.name = "SAMPLE CANDIDATE"
    as_claude_sees_it.links[1].url = ""  # GitHub
    as_claude_sees_it.links.append(type(sample.links[0])(label="Portfolio", url=""))  # no target in the PDF
    fake_api.reply(as_claude_sees_it.model_dump_json())

    profile = onboarding.extract_profile(resume)

    assert profile.name == "Sample Candidate"
    assert [(link.label, link.url) for link in profile.links] == [
        ("LinkedIn", "https://www.linkedin.com/in/sample-candidate"),
        ("GitHub", "https://github.com/sample-candidate"),  # filled from the PDF, not the project link
    ]


def test_docx_resume_becomes_text(tmp_path):
    path = tmp_path / "resume.docx"
    document = docx.Document()
    document.add_paragraph("Asha Rao")
    document.add_paragraph("Built billing APIs in Go")
    document.save(str(path))

    content = onboarding.resume_content(path)
    assert content[0]["type"] == "text"
    assert "Asha Rao" in content[0]["text"] and "Built billing APIs in Go" in content[0]["text"]


def test_render_escapes_html():
    profile = load_profile(SAMPLE_PROFILE)
    profile.experience[0].bullets[0] = "<script>alert(1)</script> & more"
    html = render_html(profile)
    assert "<script>alert" not in html
    assert "&lt;script&gt;" in html


def test_render_pdf_keeps_text_and_urls_intact(tmp_path):
    profile = load_profile(SAMPLE_PROFILE)
    out = render_pdf(profile, tmp_path / "resume.pdf")

    text = " ".join(page.extract_text() for page in PdfReader(out).pages)
    assert len(PdfReader(out).pages) == 1
    # A URL wrapped at its hyphen extracts as "samplecandidate", a broken link for ATS parsers.
    assert "github.com/sample-candidate" in text
    for bullet in profile.bullets().values():
        assert bullet.split(",")[0] in " ".join(text.split())


def test_render_pdf_never_splits_hyphenated_words(tmp_path):
    profile = load_profile(SAMPLE_PROFILE)
    words = [f"react-native-reanimated-{i}" for i in range(30)]
    profile.skills[0].skills = words
    profile.experience[0].bullets[0] = " ".join(["Resolved building-selection edge cases."] * 12)

    text = " ".join(page.extract_text() for page in PdfReader(render_pdf(profile, tmp_path / "r.pdf")).pages)
    flat = " ".join(text.split())
    assert all(word in flat for word in words)
    assert flat.count("building-selection") == 12


def test_unfilled_finds_placeholders():
    data = {"a": "TODO", "b": {"c": "x", "d": "TODO"}, "e": ["TODO"], "f": ["Pune"], "g": ""}
    assert settings.unfilled(data) == ["a", "b.d", "e"]
