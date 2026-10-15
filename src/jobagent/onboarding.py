"""Resume file (PDF / DOCX / TXT) → Profile, via one structured Claude call."""

import base64
import re
from pathlib import Path
from typing import Any

import docx
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from jobagent import llm
from jobagent.profile import Profile

SYSTEM = """You convert a candidate's resume into structured data for a job-application agent. \
The output is used to fill real job applications, so accuracy matters more than polish.

Rules:
- Copy facts exactly as written: keep each bullet's wording, numbers, names and technologies unchanged. \
Fix only obvious line-break or hyphenation artifacts from the file conversion.
- Do not invent, infer or embellish anything that is not in the resume.
- Use an empty string or empty list when something is missing.
- Keep the resume's order for roles, projects, bullets and education.
- Dates: copy as written (e.g. "Jan 2022", "2019", "Present").
- headline: the candidate's own title line if there is one, otherwise the most recent job title.
- summary: the resume's summary or objective, verbatim; empty if there is none.
- skills: keep the resume's own skill groups; if it has a flat list, use one group named "Skills". \
Include only skills that appear in the resume.
- links: LinkedIn, GitHub, portfolio and similar, with the URL as written.
- education.score: percentage or CGPA as written; empty if not given."""


def _docx_text(path: Path) -> str:
    document = docx.Document(str(path))
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(line for line in lines if line.strip())


def resume_content(path: Path) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        data = base64.standard_b64encode(path.read_bytes()).decode()
        resume_block = {
            "type": "document",
            "source": {"type": "base64", "media_type": "application/pdf", "data": data},
        }
    elif suffix == ".docx":
        resume_block = {"type": "text", "text": f"<resume>\n{_docx_text(path)}\n</resume>"}
    elif suffix in (".txt", ".md"):
        resume_block = {"type": "text", "text": f"<resume>\n{path.read_text()}\n</resume>"}
    else:
        raise ValueError(f"Unsupported resume format {suffix!r}: use .pdf, .docx, .txt or .md")
    return [resume_block, {"type": "text", "text": "Extract this resume."}]


def pdf_link_targets(path: Path) -> list[str]:
    """URLs behind clickable text ("GitHub") in a PDF: Claude sees the text, not the link target."""
    urls: list[str] = []
    try:
        pages = PdfReader(path).pages
        for page in pages:
            for annotation in page.get("/Annots") or []:
                action = annotation.get_object().get("/A") or {}
                uri = str(action.get("/URI", "")).strip()
                if uri.startswith("http") and uri not in urls:
                    urls.append(uri)
    except PdfReadError:
        return []  # unreadable link layer: Claude's text extraction still stands
    return urls


def _site(url: str) -> str:
    """"https://www.github.com/x" → "github" (compared with labels like "GitHub")."""
    host = re.sub(r"^https?://(www\.)?", "", url).split("/")[0].lower()
    return host.removesuffix(".com").removesuffix(".in")


def fill_link_targets(profile: Profile, urls: list[str]) -> None:
    """Fill empty link URLs from the PDF's link targets, matched by label ("GitHub" ↔ github.com).

    Links still without a URL are dropped (they would render as blank). Unlabelled extra URLs are not
    added: a PDF's link targets also include project and company links that don't belong in contacts.
    """
    for link in profile.links:
        if not link.url:
            link.url = next((url for url in urls if _site(url) == link.label.lower()), "")
    profile.links = [link for link in profile.links if link.url]


def extract_profile(path: Path) -> Profile:
    profile = llm.parse(
        Profile,
        task="onboarding",
        system=SYSTEM,
        content=resume_content(path),
        effort="medium",
        max_tokens=16000,
    )
    if profile.name.isupper():  # "ASHA RAO" as printed → "Asha Rao" for forms
        profile.name = profile.name.title()
    if path.suffix.lower() == ".pdf":
        fill_link_targets(profile, pdf_link_targets(path))
    return profile
